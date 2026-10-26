import re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from deck_spec import *

HEAD, BODYF, MONO = "Merriweather", "Open Sans", "Consolas"
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
BLANK = prs.slide_layouts[6]


def runs(p, text, size, color, font, bold=False):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run(); r.text = part
        f = r.font; f.size = Pt(size); f.name = font; f.bold = bold or (i % 2 == 1)
        f.color.rgb = rgb(INK if (i % 2 == 1 and color == BODY) else color)


def bullet(p, color=ACC):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.2)))); pPr.set("indent", str(-Emu(Inches(0.2))))
    bc = pPr.makeelement(qn("a:buClr"), {}); s = bc.makeelement(qn("a:srgbClr"), {"val": color}); bc.append(s); pPr.append(bc)
    bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}); pPr.append(bf)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"}); pPr.append(bu)


def textbox(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    return tf


def add_text(sl, b):
    tf = textbox(sl, b["x"], b["y"], b["w"], b["h"])
    first = True
    for para in b["paras"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = PP_ALIGN.CENTER if b["align"] == "center" else PP_ALIGN.LEFT
        if para.startswith("## "):
            p.space_before = Pt(b["size"] * 0.55); p.space_after = Pt(1)
            runs(p, para[3:], b["size"] + 0.5, NAVY, BODYF, bold=True)
        elif para.startswith("• "):
            bullet(p); p.space_after = Pt(b["size"] * 0.35)
            runs(p, para[2:], b["size"], b["color"], BODYF)
        else:
            p.space_after = Pt(b["size"] * 0.35)
            runs(p, para, b["size"], b["color"], BODYF)


def rect(sl, x, y, w, h, fill, line=None):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s


def add_img(sl, b):
    pic = sl.shapes.add_picture(b["path"], Inches(b["x"]), Inches(b["y"]), Inches(b["w"]), Inches(b["h"]))
    if b["border"]:
        pic.line.color.rgb = rgb("C9CFD6"); pic.line.width = Pt(0.75)


def add_code(sl, b):
    rect(sl, b["x"], b["y"], b["w"], b["h"], "F4F6F8", "DCE6F1")
    tf = textbox(sl, b["x"] + 0.1, b["y"] + 0.08, b["w"] - 0.2, b["h"] - 0.16)
    for i, line in enumerate(b["text"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line; r.font.name = MONO; r.font.size = Pt(b["size"])
        r.font.color.rgb = rgb(NAVY if line.rstrip().endswith("/") and not line.startswith(" ") else INK)
        if line.lstrip().startswith("#"): r.font.color.rgb = rgb(GRAY)


def add_table(sl, b):
    rows, cols = b["rows"], b["cols"]
    sh = sl.shapes.add_table(len(rows), len(cols), Inches(b["x"]), Inches(b["y"]), Inches(b["w"]), Inches(b["rowh"] * len(rows)))
    t = sh.table
    tblPr = t._tbl.tblPr; tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None: style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No style, table grid
    for j, cw in enumerate(cols): t.columns[j].width = Inches(cw)
    header = b.get("header", True) and len(rows) > 1 and not rows[0][0].startswith(("Aug", "flight"))
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(b["rowh"])
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            hdr = header and i == 0
            c.fill.solid(); c.fill.fore_color.rgb = rgb(NAVY if hdr else ("FFFFFF" if i % 2 else "F4F6F8"))
            tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
            runs(p, val, b["size"], "FFFFFF" if hdr else (NAVY if j == 0 else BODY), BODYF, bold=hdr or (j == 0 and not header and False))
            if j == 0 and not hdr:
                for r in p.runs: r.font.color.rgb = rgb(NAVY)


def add_stmt(sl, b):
    rect(sl, b["x"], b["y"], b["w"], b["h"], "F2F5F9")
    rect(sl, b["x"], b["y"], 0.07, b["h"], RED)
    tf = textbox(sl, b["x"] + 0.2, b["y"], b["w"] - 0.3, b["h"], MSO_ANCHOR.MIDDLE)
    runs(tf.paragraphs[0], b["text"], b["size"], INK, BODYF)


def add_box(sl, b):
    rect(sl, b["x"], b["y"], b["w"], b["h"], b["fill"])
    tf = textbox(sl, b["x"] + 0.12, b["y"] + 0.08, b["w"] - 0.24, b["h"] - 0.12)
    p = tf.paragraphs[0]; runs(p, b["head"], b["size"] + 1.5, b["headcolor"], BODYF, bold=True); p.space_after = Pt(3)
    for k, line in enumerate(b["body"].split("\n")):
        p = tf.add_paragraph(); runs(p, line, b["size"], BODY, BODYF)


def chrome(sl, title, n):
    tf = textbox(sl, 0.69, 0.4, 8.62, 0.66, MSO_ANCHOR.BOTTOM)
    runs(tf.paragraphs[0], title, 22 if len(title) < 52 else 19.5, RED, HEAD, bold=True)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.69), Inches(1.125), Inches(9.31), Inches(1.125))
    ln.line.color.rgb = rgb(RULE); ln.line.width = Pt(1.5)
    ft = textbox(sl, 0.69, 5.3, 8.62, 0.25)
    p = ft.paragraphs[0]; runs(p, "FAIR Data Curation  ·  ISU TrAC", 8, GRAY, BODYF)
    ft2 = textbox(sl, 8.6, 5.3, 0.71, 0.25); p = ft2.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT; runs(p, str(n), 8, GRAY, BODYF)


def title_slide(sl, s):
    rect(sl, 0, 0, 10, 0.12, RED)
    tf = textbox(sl, 0.69, 0.75, 8.62, 0.3); runs(tf.paragraphs[0], s["eyebrow"], 10.5, RULE, BODYF, bold=True)
    tf = textbox(sl, 0.69, 1.05, 8.62, 0.85, MSO_ANCHOR.BOTTOM); runs(tf.paragraphs[0], s["title"], 32, RED, HEAD, bold=True)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.69), Inches(2.02), Inches(9.31), Inches(2.02))
    ln.line.color.rgb = rgb(RULE); ln.line.width = Pt(1.5)
    tf = textbox(sl, 0.69, 2.12, 8.62, 0.7); runs(tf.paragraphs[0], s["subtitle"], 17, BODY, BODYF)
    tw, gap = 2.02, 0.18
    for i, (L_, word, sub) in enumerate(s["tiles"]):
        x = 0.69 + i * (tw + gap)
        rect(sl, x, 3.05, tw, 1.45, [TINT_B, TINT_I, TINT_G, "F2F2F2"][i])
        tf = textbox(sl, x + 0.15, 3.1, 0.7, 0.8); runs(tf.paragraphs[0], L_, 40, RED, HEAD, bold=True)
        tf = textbox(sl, x + 0.15, 3.85, tw - 0.25, 0.6)
        runs(tf.paragraphs[0], word, 13, NAVY, BODYF, bold=True)
        p = tf.add_paragraph(); runs(p, sub, 10.5, BODY, BODYF)
    tf = textbox(sl, 0.69, 4.8, 8.62, 0.45); runs(tf.paragraphs[0], s["byline"], 11, BODY, BODYF)


for n, s in enumerate(SLIDES, 1):
    sl = prs.slides.add_slide(BLANK)
    if s.get("kind") == "title":
        title_slide(sl, s)
    else:
        chrome(sl, s["title"], n)
        for b in s["blocks"]:
            {"text": add_text, "img": add_img, "code": add_code, "table": add_table, "stmt": add_stmt, "box": add_box,
             "src": lambda sl, b: runs(textbox(sl, 0.69, 5.08, 8.62, 0.22).paragraphs[0], "Source: " + b["text"], 8, GRAY, BODYF)}[b["k"]](sl, b)
    sl.notes_slide.notes_text_frame.text = s["notes"]

out = sys.argv[1] if len(sys.argv) > 1 else "FAIR_Data_Curation_Talk.pptx"
prs.save(out); print("saved", out, len(prs.slides))
